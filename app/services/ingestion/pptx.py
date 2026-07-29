from pptx import Presentation

def extract_pptx(file_path:str):
    prs=Presentation(file_path)
    slides=[]
    for slide in prs.slides:
        text=""
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text+=shape.text+"\n"
        if text.strip():
            slides.append(text)
    return slides        
    