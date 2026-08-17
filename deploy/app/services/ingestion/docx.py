from docx import Document


def extract_docx(file_path: str):
    doc = Document(file_path)

    parts = []

    # Extract paragraphs
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Extract tables
    for table_index, table in enumerate(doc.tables, start=1):

        table_rows = []

        for row in table.rows:

            cells = []

            for cell in row.cells:
                cells.append(cell.text.strip())

            if any(cells):
                table_rows.append(" | ".join(cells))

        if table_rows:

            table_text = (
                f"TABLE {table_index}\n"
                + "\n".join(table_rows)
            )

            parts.append(table_text)

    return ["\n\n".join(parts)]